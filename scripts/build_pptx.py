#!/usr/bin/env python3
"""Editable figure pack: one slide per figure in output/figures (title + image)."""
import glob, os
from pptx import Presentation
from pptx.util import Inches, Pt

FIGDIR = "output/figures"
OUT = "manuscript/figures_editable.pptx"
os.makedirs("manuscript", exist_ok=True)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
for p in sorted(glob.glob(f"{FIGDIR}/*.png")):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    name = os.path.splitext(os.path.basename(p))[0]
    tb = s.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(9.4), Inches(0.6))
    tb.text_frame.text = name.replace("_", " ")
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    s.shapes.add_picture(p, Inches(0.5), Inches(1.0), width=Inches(9.0))
prs.save(OUT)
print("wrote", OUT)
